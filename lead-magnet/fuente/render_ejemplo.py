# -*- coding: utf-8 -*-
"""Render 3D del equipo de ejemplo del Excel (para el PDF)."""
import math, os, sys
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from mpl_toolkits.mplot3d.art3d import Poly3DCollection
from openpyxl import load_workbook
AQUI = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(os.path.dirname(AQUI), 'plantilla_dynamo', 'ADLB_Plantilla_Familia.xlsx'), data_only=True)

def filas(h):
    ws = wb[h]; enc = [c.value for c in ws[1]]
    return [dict(zip(enc, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if r[0].value]

def prisma(ax, eje, a0, a1, poly, col, ec='#555'):
    idx = {'X': 0, 'Y': 1, 'Z': 2}; sec = {'X': ('Y', 'Z'), 'Y': ('X', 'Z'), 'Z': ('X', 'Y')}[eje]
    def P(a, u, v):
        q = [0, 0, 0]; q[idx[eje]] = a; q[idx[sec[0]]] = u; q[idx[sec[1]]] = v; return q
    b = [P(a0, u, v) for u, v in poly]; t = [P(a1, u, v) for u, v in poly]; n = len(poly)
    f = [b, t] + [[b[i], b[(i + 1) % n], t[(i + 1) % n], t[i]] for i in range(n)]
    CARAS.extend((cara, col, ec) for cara in f)

CARAS = []


def circ(cu, cv, r, n=40):
    return [(cu + r * math.cos(2 * math.pi * i / n), cv + r * math.sin(2 * math.pi * i / n)) for i in range(n)]

fig = plt.figure(figsize=(9, 7)); ax = fig.add_subplot(111, projection='3d', computed_zorder=False)
for p in filas('2_Piezas'):
    x0, x1, y0, y1, z0, z1 = [float(p[k]) for k in ('X_desde', 'X_hasta', 'Y_desde', 'Y_hasta', 'Z_desde', 'Z_hasta')]
    col = p.get('Color') or '#cccccc'
    if p['Forma'] == 'Cilindro':
        e = p['Eje']
        if e == 'Z': prisma(ax, 'Z', z0, z1, circ((x0 + x1) / 2, (y0 + y1) / 2, min(x1 - x0, y1 - y0) / 2), col)
        if e == 'X': prisma(ax, 'X', x0, x1, circ((y0 + y1) / 2, (z0 + z1) / 2, min(y1 - y0, z1 - z0) / 2), col)
        if e == 'Y': prisma(ax, 'Y', y0, y1, circ((x0 + x1) / 2, (z0 + z1) / 2, min(x1 - x0, z1 - z0) / 2), col)
    else:
        prisma(ax, 'Z', z0, z1, [(x0, y0), (x1, y0), (x1, y1), (x0, y1)], col)
for c in filas('3_Conexiones'):
    d = c['Direccion']; e = d[1]; s = 1 if d[0] == '+' else -1
    P = {'X': float(c['X']), 'Y': float(c['Y']), 'Z': float(c['Z'])}
    L = float(c['Largo_Boquilla'] or 150); sec = {'X': ('Y', 'Z'), 'Y': ('X', 'Z'), 'Z': ('X', 'Y')}[e]
    a1 = P[e]; a0 = a1 - s * L; cu, cv = P[sec[0]], P[sec[1]]
    if c['Tipo'] == 'Electrico':
        prisma(ax, e, min(a0, a1), max(a0, a1), [(cu - 50, cv - 50), (cu + 50, cv - 50), (cu + 50, cv + 50), (cu - 50, cv + 50)], '#E76A1F')
        continue
    dn = float(c['Diametro']); prisma(ax, e, min(a0, a1), max(a0, a1), circ(cu, cv, dn / 2 + 4), '#B0B3B6')
    if c['Brida'] == 'Si':
        od = dn * 1.2 + 105; t = 18 + dn * 0.05
        prisma(ax, e, min(a1, a1 - s * t), max(a1, a1 - s * t), circ(cu, cv, od / 2), '#E76A1F')
import numpy as np
el, az = np.radians(20), np.radians(-55)
cam = np.array([np.cos(el) * np.cos(az), np.cos(el) * np.sin(az), np.sin(el)])
for cara, col, ec in sorted(CARAS, key=lambda c: np.mean(np.array(c[0]), axis=0).dot(cam)):
    ax.add_collection3d(Poly3DCollection([cara], facecolor=col, edgecolor=ec, linewidth=0.1))
ax.set_xlim(-100, 1700); ax.set_ylim(-250, 1150); ax.set_zlim(0, 1750)
ax.set_box_aspect((1800, 1400, 1750)); ax.view_init(elev=20, azim=-55); ax.set_axis_off()
fig.patch.set_alpha(0); ax.patch.set_alpha(0)
fig.savefig(os.path.join(AQUI, 'ejemplo_3d.png'), dpi=220, transparent=True, bbox_inches='tight', pad_inches=0)
print('ok')
