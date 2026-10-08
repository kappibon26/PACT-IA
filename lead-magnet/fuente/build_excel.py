# -*- coding: utf-8 -*-
"""Genera plantilla_dynamo/ADLB_Plantilla_Familia.xlsx (con un equipo de ejemplo)."""
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = os.path.dirname(os.path.abspath(__file__))
DST = os.path.join(os.path.dirname(AQUI), 'plantilla_dynamo', 'ADLB_Plantilla_Familia.xlsx')

CARBON, NARANJA, OFF, GRIS = '0E0E0E', 'E76A1F', 'F5F4F0', 'D9D7D0'
F_ENC = Font(name='Inter', bold=True, color='FFFFFF', size=10)
F_TXT = Font(name='Inter', size=10)
F_AYUDA = Font(name='Inter', size=9, italic=True, color='6B6B66')
F_TIT = Font(name='Inter', bold=True, size=16, color=CARBON)
FILL_ENC = PatternFill('solid', fgColor=CARBON)
FILL_EJ = PatternFill('solid', fgColor='FFFFFF')
FILL_AYUDA = PatternFill('solid', fgColor=OFF)
LINEA = Border(bottom=Side(style='thin', color=GRIS))

LISTAS = {
    'Categoria': ['Equipo mecanico', 'Equipo electrico', 'Aparato sanitario', 'Accesorio de tuberia', 'Modelo generico'],
    'Forma': ['Caja', 'Cilindro'],
    'Eje': ['X', 'Y', 'Z'],
    'Detalle': ['Siempre', 'Medio y fino'],
    'TipoConexion': ['Tuberia', 'Ducto', 'Electrico'],
    'FormaConexion': ['Redonda', 'Rectangular'],
    'Direccion': ['+X', '-X', '+Y', '-Y', '+Z', '-Z'],
    'SiNo': ['Si', 'No'],
    'Sistema': ['Agua fria', 'Agua caliente', 'Sanitario', 'Ventilacion', 'Pluvial', 'Incendio',
                'Suministro hidronico', 'Retorno hidronico', 'Suministro aire', 'Retorno aire', 'Extraccion', 'Otro'],
    'Flujo': ['Entra', 'Sale', 'Ambos'],
    'TipoDato': ['Texto', 'Longitud', 'Numero', 'Entero', 'SiNo', 'Angulo', 'Area', 'Volumen', 'Material', 'URL'],
    'Grupo': ['Identidad', 'Datos', 'Dimensiones', 'Restricciones', 'Materiales', 'Mecanico', 'Electrico',
              'Fontaneria', 'Graficos', 'Texto', 'Estructural', 'General'],
}


def hoja_tabla(wb, nombre, columnas, filas, anchos, validaciones, titulo):
    ws = wb.create_sheet(nombre)
    for i, (col, ayuda) in enumerate(columnas, 1):
        c = ws.cell(row=1, column=i, value=col)
        c.font, c.fill = F_ENC, FILL_ENC
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.comment = Comment(ayuda, 'ADLB')
        ws.column_dimensions[c.column_letter].width = anchos[i - 1]
    ws.row_dimensions[1].height = 30
    for r, fila in enumerate(filas, 2):
        for i, v in enumerate(fila, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.fill, c.border = F_TXT, FILL_EJ, LINEA
    ws.freeze_panes = 'B2'
    for letra, lista in validaciones:
        dv = DataValidation(type='list', formula1='=Listas!$%s$2:$%s$%d' % (
            LISTA_COL[lista], LISTA_COL[lista], len(LISTAS[lista]) + 1), allow_blank=True)
        dv.error, dv.errorTitle = u'Elige un valor de la lista', u'ADLB'
        ws.add_data_validation(dv)
        dv.add('%s2:%s300' % (letra, letra))
    ws.sheet_properties.tabColor = NARANJA
    ws.cell(row=1, column=len(columnas) + 2, value=titulo).font = F_AYUDA
    return ws


wb = Workbook()
leeme = wb.active
leeme.title = 'LEEME'
leeme.sheet_view.showGridLines = False
leeme.column_dimensions['A'].width = 4
leeme.column_dimensions['B'].width = 110
lineas = [
    ('ADLB. Generador de familias de Revit', F_TIT),
    ('Llena las 4 hojas naranjas y corre el grafico ADLB_Generar_Familia.dyn en Dynamo.', F_TXT),
    ('', F_TXT),
    ('ANTES DE EMPEZAR: prepara tu dibujo', Font(name='Inter', bold=True, size=11, color=NARANJA)),
    ('- Dibujo en milimetros. Planta y alzado frontal del equipo (DWG o DXF).', F_TXT),
    ('- Elige un ORIGEN facil de ubicar en obra (esquina de la base o eje del equipo) y llevalo al 0,0 del dibujo.', F_TXT),
    ('- En el alzado, el PISO debe quedar en Y = 0 del dibujo. Asi Z = altura real.', F_TXT),
    ('- X = largo en la planta (izquierda a derecha). Y = ancho en la planta (abajo hacia arriba). Z = altura.', F_TXT),
    ('', F_TXT),
    ('COMO LLENAR CADA HOJA', Font(name='Inter', bold=True, size=11, color=NARANJA)),
    ('1_Familia      Nombre del archivo, categoria, tipo y prefijo de tus parametros.', F_TXT),
    ('2_Piezas       Una fila por pieza. Lee en la PLANTA de donde a donde va en X y en Y; en el ALZADO, en Z.', F_TXT),
    ('               Forma Caja o Cilindro. Un cilindro necesita su Eje (Z si es vertical).', F_TXT),
    ('3_Conexiones   Una fila por conexion. X, Y, Z = centro de la CARA de la brida. Direccion = hacia donde sale el tubo.', F_TXT),
    ('4_Parametros   Tus datos (codigo, peso, capacidad...). Usa la convencion PREFIJO_Concepto_Unidad.', F_TXT),
    ('', F_TXT),
    ('REGLAS PARA NOMBRAR PARAMETROS (sirven en cualquier oficina)', Font(name='Inter', bold=True, size=11, color=NARANJA)),
    ('- PREFIJO_Concepto_Unidad   ->  EQ_Peso_kg, EQ_Capacidad_m3h, EQ_Codigo', F_TXT),
    ('- Sin espacios, tildes, enes ni simbolos. Palabras con mayuscula inicial unidas por guion bajo.', F_TXT),
    ('- La unidad va al final solo si el tipo de dato no la trae (Numero o Texto). Longitud ya va en mm.', F_TXT),
    ('- Si/No: el concepto empieza con Mostrar_ o Es_ (EQ_Mostrar_Zona_Servicio).', F_TXT),
    ('- Instancia = cambia en cada equipo colocado (codigo, sistema). Tipo = es del modelo (peso, capacidad).', F_TXT),
    ('- Las conexiones crean sus propios parametros: PREFIJO_Conexion_DN, _Largo, _Ancho, _Alto, _Tension...', F_TXT),
    ('', F_TXT),
    ('Pasa el mouse sobre cada encabezado para ver su ayuda. Las filas que empiezan con # se ignoran.', F_AYUDA),
    ('Escribe HOY en un parametro de texto para poner la fecha del dia.', F_AYUDA),
]
for i, (t, f) in enumerate(lineas, 2):
    leeme.cell(row=i, column=2, value=t).font = f

# --- listas (hoja oculta)
listas = wb.create_sheet('Listas')
LISTA_COL = {}
for i, (k, vals) in enumerate(LISTAS.items(), 1):
    c = listas.cell(row=1, column=i, value=k)
    LISTA_COL[k] = c.column_letter
    for j, v in enumerate(vals, 2):
        listas.cell(row=j, column=i, value=v)
listas.sheet_state = 'hidden'

# --- 1_Familia
fam = wb.create_sheet('1_Familia', 1)
fam.sheet_properties.tabColor = NARANJA
for i, (h, w) in enumerate((('Campo', 24), ('Valor', 46), ('Ayuda', 70)), 1):
    c = fam.cell(row=1, column=i, value=h)
    c.font, c.fill = F_ENC, FILL_ENC
    fam.column_dimensions[c.column_letter].width = w
datos = [
    ('Nombre_Archivo', 'ADLB_Equipo_Bombeo_Tanque', 'Nombre del .rfa. Sin espacios ni tildes.'),
    ('Categoria', 'Equipo mecanico', 'Categoria de Revit (lista).'),
    ('Nombre_Tipo', 'Tanque 500 L + bomba 7.5 kW', 'Nombre del tipo dentro de la familia.'),
    ('Prefijo', 'EQ', 'Prefijo de TUS parametros (2-4 letras): EQ, ABC, OF1...'),
    ('Fabricante', 'Fabricante de ejemplo', 'Parametro nativo de Revit.'),
    ('Modelo', 'TB-500', 'Parametro nativo de Revit.'),
    ('Descripcion', 'Tanque de almacenamiento con bomba centrifuga sobre base', 'Parametro nativo de Revit.'),
    ('URL', '', 'Ficha tecnica (opcional).'),
    ('Archivo_Compartidos', '', 'Ruta del .txt de parametros compartidos (opcional).'),
    ('Plantilla_RFT', '', 'Ruta de una plantilla .rft propia (opcional; si no, se busca sola).'),
]
for r, fila in enumerate(datos, 2):
    for i, v in enumerate(fila, 1):
        c = fam.cell(row=r, column=i, value=v)
        c.font = F_AYUDA if i == 3 else F_TXT
        c.fill = FILL_AYUDA if i != 2 else FILL_EJ
        c.border = LINEA
dv = DataValidation(type='list', formula1='=Listas!$A$2:$A$6', allow_blank=False)
fam.add_data_validation(dv)
dv.add('B3')

# --- 2_Piezas
hoja_tabla(wb, '2_Piezas', [
    ('Pieza', 'Nombre corto de la pieza (solo para ti).'),
    ('Forma', 'Caja o Cilindro.'),
    ('Eje', 'Solo cilindros: Z = vertical, X o Y = horizontal.'),
    ('X_desde', 'Planta: donde empieza la pieza en X (mm desde el origen).'),
    ('X_hasta', 'Planta: donde termina en X.'),
    ('Y_desde', 'Planta: donde empieza en Y.'),
    ('Y_hasta', 'Planta: donde termina en Y.'),
    ('Z_desde', 'Alzado: altura de la parte de abajo (mm desde el piso).'),
    ('Z_hasta', 'Alzado: altura de la parte de arriba.'),
    ('Grupo', 'Subcategoria: Cuerpo, Motor, Estructura... (para filtros y visibilidad).'),
    ('Material', 'Nombre del material. Se crea si no existe y queda como parametro Mat_<nombre>.'),
    ('Color', 'Color del material en hexadecimal, ej. #1E1E1E (opcional).'),
    ('Detalle', 'Siempre, o Medio y fino para piezas pequenas (se ocultan en nivel Bajo).'),
], [
    ['Base', 'Caja', '', 0, 1600, 0, 900, 0, 120, 'Estructura', 'Acero negro', '#1E1E1E', 'Siempre'],
    ['Tanque', 'Cilindro', 'Z', 100, 800, 100, 800, 120, 1520, 'Cuerpo', 'Acero inoxidable', '#C8CCCF', 'Siempre'],
    ['Tapa tanque', 'Cilindro', 'Z', 80, 820, 80, 820, 1520, 1550, 'Cuerpo', 'Acero inoxidable', '#C8CCCF', 'Siempre'],
    ['Soporte bomba', 'Caja', '', 900, 1550, 300, 600, 120, 200, 'Estructura', 'Acero negro', '#1E1E1E', 'Siempre'],
    ['Bomba', 'Cilindro', 'X', 950, 1150, 250, 650, 200, 600, 'Cuerpo', 'Hierro fundido', '#3B5B8C', 'Siempre'],
    ['Motor', 'Cilindro', 'X', 1150, 1550, 300, 600, 250, 550, 'Motor', 'Motor', '#2F3E4E', 'Siempre'],
    ['Caja de bornes', 'Caja', '', 1300, 1420, 600, 660, 380, 500, 'Motor', 'Motor', '#2F3E4E', 'Medio y fino'],
], [10, 10, 7, 10, 10, 10, 10, 10, 10, 14, 18, 11, 14],
    [('B', 'Forma'), ('C', 'Eje'), ('M', 'Detalle')], 'Medidas en mm. Lee X/Y en la planta y Z en el alzado.')

# --- 3_Conexiones
hoja_tabla(wb, '3_Conexiones', [
    ('Conexion', 'Nombre corto sin espacios: Entrada, Descarga, Drenaje, Motor...'),
    ('Tipo', 'Tuberia, Ducto o Electrico.'),
    ('Forma', 'Redonda o Rectangular (rectangular solo para ductos).'),
    ('X', 'Centro de la CARA de la brida (donde se conecta el tubo), en mm.'),
    ('Y', 'Centro de la cara, en mm.'),
    ('Z', 'Centro de la cara, en mm (altura).'),
    ('Direccion', 'Hacia donde sale el tubo desde el equipo: +X, -X, +Y, -Y, +Z (arriba) o -Z (abajo).'),
    ('Diametro', 'Diametro nominal en mm (tuberia o ducto redondo).'),
    ('Ancho', 'Ducto rectangular: ancho en mm.'),
    ('Alto', 'Ducto rectangular: alto en mm.'),
    ('Largo_Boquilla', 'Cuanto sobresale la boquilla del equipo, en mm (por defecto 150).'),
    ('Brida', 'Si = brida con pernos; No = extremo liso.'),
    ('Sistema', 'Clasificacion de sistema de Revit (lista).'),
    ('Flujo', 'Entra, Sale o Ambos.'),
    ('Tension_V', 'Solo electricas: tension en voltios.'),
    ('Potencia_kW', 'Solo electricas: potencia en kW.'),
], [
    ['Llenado', 'Tuberia', 'Redonda', 450, 450, 1650, '+Z', 80, '', '', 100, 'Si', 'Agua fria', 'Entra', '', ''],
    ['Descarga', 'Tuberia', 'Redonda', 1050, 450, 750, '+Z', 65, '', '', 150, 'Si', 'Otro', 'Sale', '', ''],
    ['Drenaje', 'Tuberia', 'Redonda', 450, 0, 250, '-Y', 50, '', '', 100, 'No', 'Sanitario', 'Sale', '', ''],
    ['Motor', 'Electrico', '', 1360, 700, 440, '+Y', '', '', '', 40, '', '', '', 460, 7.5],
], [12, 11, 12, 8, 8, 8, 10, 10, 8, 8, 14, 8, 16, 9, 10, 11],
    [('B', 'TipoConexion'), ('C', 'FormaConexion'), ('G', 'Direccion'), ('L', 'SiNo'), ('M', 'Sistema'), ('N', 'Flujo')],
    'X, Y, Z = centro de la cara de conexion.')

# --- 4_Parametros
hoja_tabla(wb, '4_Parametros', [
    ('Nombre', 'PREFIJO_Concepto_Unidad, sin espacios ni tildes. Ej: EQ_Peso_kg'),
    ('TipoDato', 'Texto, Longitud (mm), Numero, Entero, SiNo, Angulo, Area, Volumen, Material, URL.'),
    ('Grupo', 'Grupo de la ventana de propiedades.'),
    ('Instancia', 'Si = cambia en cada equipo colocado. No = es del tipo.'),
    ('Compartido', 'Si = se toma del archivo de compartidos (para etiquetas y tablas).'),
    ('Valor', 'Valor para todos los tipos. Longitud en mm. HOY = fecha de hoy.'),
    ('Formula', 'Formula de Revit (opcional). Si hay formula se ignora Valor.'),
], [
    ['EQ_Codigo', 'Texto', 'Identidad', 'Si', 'No', 'B-01', ''],
    ['EQ_Sistema', 'Texto', 'Identidad', 'Si', 'No', 'Agua', ''],
    ['EQ_Peso_kg', 'Numero', 'Datos', 'No', 'No', 650, ''],
    ['EQ_Capacidad_m3h', 'Numero', 'Datos', 'No', 'No', 10, ''],
    ['EQ_Volumen_Tanque_L', 'Numero', 'Datos', 'No', 'No', 500, ''],
    ['EQ_LOD', 'Texto', 'Datos', 'No', 'No', '350', ''],
    ['EQ_Fuente_Geometria', 'Texto', 'Datos', 'No', 'No', 'Plano del fabricante', ''],
    ['EQ_Fecha_Modelado', 'Texto', 'Datos', 'No', 'No', 'HOY', ''],
    ['EQ_Holgura_Mantenimiento', 'Longitud', 'Restricciones', 'Si', 'No', 600, ''],
    ['EQ_Mostrar_Zona_Servicio', 'SiNo', 'Graficos', 'Si', 'No', 'Si', ''],
], [28, 12, 15, 10, 12, 22, 22],
    [('B', 'TipoDato'), ('C', 'Grupo'), ('D', 'SiNo'), ('E', 'SiNo')], 'Convencion: PREFIJO_Concepto_Unidad.')

wb.move_sheet('Listas', offset=10)
wb.save(DST)
print('ok ->', DST)
